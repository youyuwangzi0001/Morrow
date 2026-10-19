#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""M09 校验：XLSX 是真实 OOXML、指定单元格与真公式正确，并独立重算 D4。

用法：python check_xlsx.py <xlsx>
输出 JSON {ok, detail, sheet, cells, formulas, recalculated}
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "detail": "usage: check_xlsx.py <xlsx>"}, ensure_ascii=False))
        return 2
    path = Path(sys.argv[1])
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "detail": f"openpyxl 不可用：{exc}"}, ensure_ascii=False))
        return 1
    if not path.exists() or path.stat().st_size == 0:
        print(json.dumps({"ok": False, "detail": "文件不存在或为空"}, ensure_ascii=False))
        return 1
    try:
        # 必须能按 OOXML 打开（不是截图/改后缀）
        wb = load_workbook(path, data_only=False)
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "detail": f"无法作为 xlsx 打开：{exc}"}, ensure_ascii=False))
        return 1
    sheet_name = "Sales" if "Sales" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    cells = {f"{col}{row}": ws[f"{col}{row}"].value for row in range(1, 5) for col in "ABCD"}
    formulas = {
        "D2": ws["D2"].value,
        "D3": ws["D3"].value,
        "D4": ws["D4"].value,
    }
    # 独立重算：只在公式语义符合题目要求时计算（不信任缓存值）
    def eval_formula(formula, values):
        if not isinstance(formula, str) or not formula.startswith("="):
            return None
        expr = formula[1:].replace(" ", "").upper()
        b2, c2 = values.get("B2"), values.get("C2")
        b3, c3 = values.get("B3"), values.get("C3")
        try:
            if expr in ("B2*C2",):
                return b2 * c2
            if expr in ("B3*C3",):
                return b3 * c3
            if expr in ("SUM(D2:D3)", "D2+D3", "SUM(D2,D3)"):
                d2, d3 = values.get("D2"), values.get("D3")
                return d2 + d3
        except Exception:  # noqa: BLE001
            return None
        return None

    problems = []
    if cells.get("A1") != "Item" or cells.get("B1") != "Qty" or cells.get("C1") != "Price":
        problems.append(f"A1:C1 表头不符：{cells.get('A1')}/{cells.get('B1')}/{cells.get('C1')}")
    if cells.get("A2") != "Pen" or cells.get("B2") != 2 or cells.get("C2") != 3:
        problems.append(f"第 2 行数据不符：{cells.get('A2')}/{cells.get('B2')}/{cells.get('C2')}")
    if cells.get("A3") != "Book" or cells.get("B3") != 3 or cells.get("C3") != 5:
        problems.append(f"第 3 行数据不符：{cells.get('A3')}/{cells.get('B3')}/{cells.get('C3')}")
    numeric = {}
    for key, value in cells.items():
        if isinstance(value, (int, float)):
            numeric[key] = value
    recalculated = {}
    for cell in ("D2", "D3"):
        result = eval_formula(formulas.get(cell), numeric)
        recalculated[cell] = result
        if result is None:
            problems.append(f"{cell} 不是题目要求的真公式（实际：{formulas.get(cell)!r}）")
    if isinstance(recalculated.get("D2"), (int, float)) and recalculated["D2"] != 6:
        problems.append(f"D2 重算期望 6，实际 {recalculated['D2']}")
    if isinstance(recalculated.get("D3"), (int, float)) and recalculated["D3"] != 15:
        problems.append(f"D3 重算期望 15，实际 {recalculated['D3']}")
    # D4 依赖 D2/D3 的公式结果：用重算值再算 SUM
    values = dict(numeric)
    if isinstance(recalculated.get("D2"), (int, float)):
        values["D2"] = recalculated["D2"]
    if isinstance(recalculated.get("D3"), (int, float)):
        values["D3"] = recalculated["D3"]
    d4 = eval_formula(formulas.get("D4"), values)
    if d4 is None:
        problems.append(f"D4 不是题目要求的真公式（实际：{formulas.get('D4')!r}）")
    if isinstance(d4, (int, float)) and d4 != 21:
        problems.append(f"D4 重算期望 21，实际 {d4}")
    recalculated["D4"] = d4
    soffice_info = {'available': False}
    soffice_path = os.environ.get('MORROW_SOFFICE') or r'C:\Program Files\LibreOffice\program\soffice.exe'
    if os.path.exists(soffice_path):
        soffice_info['available'] = True
        tmpdir = Path(tempfile.mkdtemp(prefix='morrow-xlsx-recalc-'))
        try:
            outdir = tmpdir / 'out'
            outdir.mkdir(parents=True, exist_ok=True)
            proc = subprocess.run(
                [soffice_path, '--headless', '--convert-to', 'xlsx', '--outdir', str(outdir), str(path)],
                capture_output=True,
                timeout=120,
            )
            converted = outdir / path.name
            if proc.returncode == 0 and converted.exists():
                wb2 = load_workbook(converted, data_only=True)
                ws2 = wb2[sheet_name] if sheet_name in wb2.sheetnames else wb2[wb2.sheetnames[0]]
                d2, d3, d4 = ws2['D2'].value, ws2['D3'].value, ws2['D4'].value
                soffice_info.update({'ok': d2 == 6 and d3 == 15 and d4 == 21, 'd2': d2, 'd3': d3, 'd4': d4})
                if not soffice_info['ok']:
                    problems.append(f'soffice 重算不符：D2={d2}, D3={d3}, D4={d4}')
            else:
                soffice_info['error'] = (proc.stderr or b'').decode('utf-8', 'ignore')[:300]
        except Exception as exc:  # noqa: BLE001
            soffice_info['error'] = f'{type(exc).__name__}: {exc}'
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)
    ok = not problems
    print(
        json.dumps(
            {
                "ok": ok,
                "detail": "；".join(problems) if problems else "OOXML/表头/数据/真公式/重算 21 均正确",
                "sheet": sheet_name,
                "cells": {key: (str(value) if not isinstance(value, (int, float)) else value) for key, value in cells.items()},
                "formulas": formulas,
                "recalculated": recalculated,
            },
            ensure_ascii=False,
            default=str,
        )
    )
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
