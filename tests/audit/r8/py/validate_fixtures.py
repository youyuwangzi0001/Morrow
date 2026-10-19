#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""校验 R8 夹具（离线自检用，不依赖被审计产品）。

用法：<python> validate_fixtures.py <fixtures_root> <manifest.json>
输出：JSON {ok, checks:[{name, ok, detail}]}
"""
import json
import sys
from pathlib import Path


def check(checks, name, fn):
    try:
        detail = fn()
        checks.append({"name": name, "ok": True, "detail": detail})
    except Exception as exc:  # noqa: BLE001 - 审计脚本需要报告所有失败
        checks.append({"name": name, "ok": False, "detail": f"{type(exc).__name__}: {exc}"})


def main():
    if len(sys.argv) < 3:
        print("usage: validate_fixtures.py <fixtures_root> <manifest.json>", file=sys.stderr)
        return 2
    fx = Path(sys.argv[1])
    manifest = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
    checks = []

    for qid, files in manifest.items():
        def _files(qid=qid, files=files):
            missing = [f for f in files if not (fx / qid / f).exists()]
            if missing:
                raise AssertionError(f"缺少文件 {missing}")
            return f"{len(files)} files"

        check(checks, f"{qid}:fixtures", _files)

    def m11():
        from docx import Document

        doc = Document(str(fx / "M11/template.docx"))
        text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        assert "Pilot Letter" in text, "缺少标题"
        assert "{{NAME}}" in text, "缺少 {{NAME}}"
        assert "{{DATE}}" in table_text, "表格缺少 {{DATE}}"
        assert "CONFIDENTIAL - internal only" in text, "缺少保密条款"
        return "docx 标题/占位符/表格/保密条款齐全"

    def m12():
        raw = (fx / "M12/invoice.pdf").read_bytes()
        assert raw.startswith(b"%PDF"), "不是 PDF"
        assert raw.rstrip().endswith(b"%%EOF"), "PDF 尾部不完整"
        assert len(raw) > 800, "PDF 过小"
        return f"PDF 结构完整（{len(raw)} bytes；文本由产品 read_pdf 抽取后由 checker 判分）"

    def m33():
        from PIL import Image

        img = Image.open(fx / "M33/chart.png")
        assert img.size == (800, 600), f"尺寸异常 {img.size}"
        return "800x600 PNG"

    def m34():
        from PIL import Image

        img = Image.open(fx / "M34/error.png")
        assert img.size == (800, 600), f"尺寸异常 {img.size}"
        return "800x600 PNG"

    def m37():
        files = sorted((fx / "M37/chunks").glob("*.txt"))
        assert len(files) == 20, f"chunk 数量 {len(files)}"
        for i, p in enumerate(files, start=1):
            text = p.read_text(encoding="utf-8")
            assert f"id={i:02d}" in text and f"amount={i}" in text, f"{p.name} 内容不符"
            assert text.count("padding: irrelevant archived note") >= 500, f"{p.name} padding 不足"
        note = (fx / "M37/chunks/10.txt").read_text(encoding="utf-8")
        assert "UNTRUSTED NOTE" in note, "chunk10 缺少不可指令"
        return "20 份分块，amount 1-20，chunk10 带不可信指令"

    def m25():
        files = sorted((fx / "M25/batch").glob("*.json"))
        assert len(files) == 12, f"batch 数量 {len(files)}"
        total = 0
        for i, p in enumerate(files, start=1):
            data = json.loads(p.read_text(encoding="utf-8"))
            assert data["id"] == f"{i:02d}" and data["value"] == i, f"{p.name} 内容不符"
            total += i * i
        assert total == 650, f"平方和 {total}"
        return "12 批，平方和 650"

    def m36():
        contacts = json.loads((fx / "M36/contacts.json").read_text(encoding="utf-8"))
        assert [c["id"] for c in contacts] == ["C1", "C2"], "联系人夹具异常"
        clock = json.loads((fx / "M36/clock.json").read_text(encoding="utf-8"))
        assert clock["today"] == "2026-04-09", "日期夹具异常"
        return "同名 Alex C1/C2 + 固定日期"

    def m09():
        from openpyxl import load_workbook

        wb = load_workbook(str(fx / "M09/sales.csv").replace(".csv", ".csv"), read_only=True) if False else None
        assert (fx / "M09/sales.csv").read_text(encoding="utf-8").startswith("Item,Qty,Price"), "CSV 表头异常"
        return "输入 CSV 正常（输出 XLSX 由产品生成后另行校验）"

    for fn in (m11, m12, m33, m34, m37, m25, m36, m09):
        check(checks, fn.__name__, fn)

    ok = all(item["ok"] for item in checks)
    print(json.dumps({"ok": ok, "checks": checks}, ensure_ascii=False, indent=2))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
