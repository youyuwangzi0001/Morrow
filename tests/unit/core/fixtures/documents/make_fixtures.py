"""Generate the real binary fixtures used by Morrow's document-tool unit tests.

Why committed binaries: the document tools must be tested against files produced by
*independent* implementations (python-docx / openpyxl / LibreOffice), not by the same
writer code that lives in src/core/tools. The generated files are checked in so
`npm run test:unit` is reproducible on any machine without Python.

Provenance / regeneration (Windows, bundled Python via load_workspace_dependencies):
    <python> make_fixtures.py
Requirements: python-docx 1.2.0, openpyxl 3.1.5, LibreOffice Kit CLI (only for PDF).

Outputs (all in this directory):
    template.docx     M11-style letter: Title + body {{NAME}} + 1x2 table {{DATE}} + confidentiality
    split-run.docx    same idea but placeholders are split across several w:r runs (robustness)
    invoice.pdf       M12-style single-page invoice: INV-0042 / 12, 18, 7 / printed total 38
    invoice.docx      source document invoice.pdf was exported from (provenance)
    input.xlsx        read fixture: Sales sheet with real formulas + a second sheet
    manifest.json     sha256 + source of every fixture
"""
from __future__ import annotations

import hashlib
import json
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
BUILD = HERE / "build"
BUILD.mkdir(exist_ok=True)

NODE = r"C:\Users\Administrator\AppData\Local\Programs\DeepSeek Harness\resources\runtime\primary-runtime\dependencies\node\bin\node.exe"
LO_CLI = r"C:\Users\Administrator\AppData\Local\Programs\DeepSeek Harness\resources\app.asar.unpacked\dsh\node_modules\@deepseek-ai\libreoffice-kit\lib\cli.js"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


# ── M11: DOCX template ────────────────────────────────────────────────
def build_template() -> Path:
    """Title / body placeholder / 1x2 table with date placeholder / confidentiality clause."""
    document = Document()
    document.add_heading("Pilot Letter", level=0)  # Title style
    document.add_paragraph("Dear {{NAME}}")
    document.add_paragraph("We are glad to confirm your pilot participation.")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Date"
    table.cell(0, 1).text = "{{DATE}}"
    document.add_paragraph("CONFIDENTIAL - internal only")
    path = HERE / "template.docx"
    document.save(path)
    return path


# ── M11 robustness: placeholders split across runs ────────────────────
def build_split_run() -> Path:
    """A real Word phenomenon: `{{NAME}}` and `{{DATE}}` split over multiple w:r runs.

    Word does this whenever spell-check / formatting boundaries fall inside a
    placeholder. A naive per-<w:t> replacement misses them, so edit_docx must
    merge the text of consecutive runs before matching.
    """
    document = Document()
    document.add_heading("Split Run Letter", level=0)
    paragraph = document.add_paragraph()
    paragraph.add_run("Dear {{NA")
    paragraph.add_run("ME")
    paragraph.add_run("}},")
    paragraph.add_run(" welcome")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Date"
    paragraph = table.cell(0, 1).paragraphs[0]
    paragraph.add_run("{{DA")
    paragraph.add_run("TE}}")
    document.add_paragraph("CONFIDENTIAL - internal only")
    path = HERE / "split-run.docx"
    document.save(path)
    return path


# ── M12: invoice document → PDF via LibreOffice ───────────────────────
def build_invoice_docx() -> Path:
    document = Document()
    document.add_heading("INVOICE", level=0)
    document.add_paragraph("Invoice ID: INV-0042")
    document.add_paragraph("Customer: Aurora Trading Ltd.")
    document.add_paragraph("Items: 12, 18, 7")
    document.add_paragraph("Printed total: 38")
    document.add_paragraph("Payment terms: net 30 days.")
    path = HERE / "invoice.docx"
    document.save(path)
    return path


def build_invoice_zh_docx() -> Path:
    """中文发票：验证 pdf-parse（pdf.js 1.10）对中文导出 PDF 的文本抽取（ToUnicode）。"""
    document = Document()
    document.add_heading("发票", level=0)
    document.add_paragraph("发票编号：INV-0042")
    document.add_paragraph("项目金额：12、18、7")
    document.add_paragraph("印刷总额：38")
    document.add_paragraph("备注：本发票为离线测试样例。")
    path = HERE / "invoice-zh.docx"
    document.save(path)
    return path


def build_blank_docx() -> Path:
    """A valid document with no text at all → exported PDF has no extractable text."""
    document = Document()
    document.add_paragraph("")
    path = BUILD / "blank.docx"
    document.save(path)
    return path


def export_pdf(docx_path: Path, pdf_path: Path) -> Path:
    # LibreOffice Kit 拒绝覆盖已存在的输出文件：先删旧产物再导出
    if pdf_path.exists():
        pdf_path.unlink()
    result = subprocess.run(
        [NODE, LO_CLI, "convert", "--input", str(docx_path), "--output", str(pdf_path)],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0 or not pdf_path.exists():
        raise SystemExit(f"LibreOffice export failed: {result.returncode}\n{result.stdout}\n{result.stderr}")
    signature = pdf_path.read_bytes()[:5]
    if signature != b"%PDF-":
        raise SystemExit(f"not a PDF: {pdf_path} {signature!r}")
    return pdf_path


# ── XLSX read fixture (openpyxl, formula cells have no cached value) ──
def build_input_xlsx() -> Path:
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    rows = [
        ("Item", "Qty", "Price", "Total"),
        ("Pen", 2, 3, "=B2*C2"),
        ("Book", 3, 5, "=B3*C3"),
        ("", "", "", "=SUM(D2:D3)"),
    ]
    for row in rows:
        sales.append(row)
    for cell in sales[1]:
        cell.font = Font(bold=True)
    notes = workbook.create_sheet("Notes")
    notes.append(("key", "value"))
    notes.append(("owner", "林宁"))
    notes.append(("approved", True))
    notes.append(("checked", 2026))
    path = HERE / "input.xlsx"
    workbook.save(path)
    return path


def main() -> None:
    template = build_template()
    split_run = build_split_run()
    invoice_docx = build_invoice_docx()
    invoice_zh_docx = build_invoice_zh_docx()
    blank_docx = build_blank_docx()
    input_xlsx = build_input_xlsx()
    invoice_pdf = export_pdf(invoice_docx, HERE / "invoice.pdf")
    invoice_zh_pdf = export_pdf(invoice_zh_docx, HERE / "invoice-zh.pdf")
    blank_pdf = export_pdf(blank_docx, HERE / "blank.pdf")

    manifest = {
        "generator": "make_fixtures.py",
        "libraries": {"python-docx": "1.2.0", "openpyxl": "3.1.5", "pdf-export": "LibreOffice Kit"},
        "fixtures": {
            path.name: {"sha256": sha256(path), "bytes": path.stat().st_size}
            for path in (template, split_run, invoice_docx, invoice_pdf, invoice_zh_docx, invoice_zh_pdf, blank_pdf, input_xlsx)
        },
        "scenarios": {
            "M11": "template.docx — Title + Dear {{NAME}} + 1x2 table {{DATE}} + CONFIDENTIAL",
            "M11-robustness": "split-run.docx — placeholders split across w:r runs",
            "M12": "invoice.pdf — INV-0042 / Items 12,18,7 / Printed total 38 (12+18+7=37)",
            "M12-中文": "invoice-zh.pdf — 中文发票：发票编号 INV-0042 / 项目金额 12、18、7 / 印刷总额 38",
            "read_xlsx": "input.xlsx — Sales sheet with real formulas =B2*C2 / =B3*C3 / =SUM(D2:D3)",
            "error-path": "blank.pdf — valid PDF with no extractable text",
        },
    }
    (HERE / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for name, info in manifest["fixtures"].items():
        print(f"{name:18s} {info['bytes']:>7d} B  {info['sha256'][:16]}…")
    print("manifest.json written")


if __name__ == "__main__":
    main()
